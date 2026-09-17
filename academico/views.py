from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.utils import timezone
from django.db import transaction
from django.db.models import Q
from django.http import JsonResponse, HttpResponse
from decimal import Decimal
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

from .models import Curso, Cursante, Inscripcion, TituloCursante
from .services.syric_parser import import_syric_file
from .services.moodle_generator import generate_moodle_csv_response

@login_required
def importar_syric(request):
    cursos = Curso.objects.filter(activo=True)
    stats = None

    if request.method == 'POST':
        curso_id = request.POST.get('curso')
        archivo = request.FILES.get('archivo')

        if not curso_id:
            messages.error(request, 'Debe seleccionar un curso de la lista.')
        elif not archivo:
            messages.error(request, 'Debe seleccionar un archivo para importar.')
        else:
            filename = archivo.name.lower()
            if not filename.endswith(('.xlsx', '.xls', '.csv')):
                messages.error(request, 'Formato de archivo inválido. Solo se admiten archivos .xlsx, .xls y .csv.')
            else:
                try:
                    stats = import_syric_file(archivo, archivo.name, curso_id)
                    if stats['errores'] and len(stats['errores']) == stats['total']:
                        messages.warning(request, 'La importación falló por completo o no contenía registros válidos.')
                    elif stats['errores']:
                        messages.warning(request, f'Importación completada con algunos errores ({len(stats["errores"])} filas fallidas).')
                    else:
                        messages.success(request, 'Importación finalizada con éxito.')
                except Exception as e:
                    messages.error(request, f'Ocurrió un error inesperado al procesar el archivo: {str(e)}')

    context = {
        'cursos': cursos,
        'stats': stats,
        'active_tab': 'importar',
    }
    return render(request, 'academico/importar_syric.html', context)


@login_required
def lista_inscripciones(request):
    cursos = Curso.objects.filter(activo=True)
    
    # Query parameters / Filters
    curso_id = request.GET.get('curso', '')
    estado_pago = request.GET.get('estado_pago', '')
    estado_moodle = request.GET.get('estado_moodle', '')
    search_query = request.GET.get('q', '')

    inscripciones = Inscripcion.objects.all().select_related('cursante', 'curso')

    # Apply filters
    if curso_id:
        inscripciones = inscripciones.filter(curso_id=curso_id)
    if estado_pago:
        inscripciones = inscripciones.filter(estado_pago=estado_pago)
    if estado_moodle:
        inscripciones = inscripciones.filter(estado_moodle=estado_moodle)
    if search_query:
        inscripciones = inscripciones.filter(
            Q(cursante__dni__icontains=search_query) |
            Q(cursante__apellido__icontains=search_query) |
            Q(cursante__nombre__icontains=search_query)
        )

    # Count calculations (relative to selected curso if filtered, otherwise global)
    base_counts = Inscripcion.objects.all()
    if curso_id:
        base_counts = base_counts.filter(curso_id=curso_id)

    total_count = base_counts.count()
    pagados_count = base_counts.filter(estado_pago__in=Inscripcion.ESTADOS_PAGO_MATRICULABLES).count()
    pendientes_count = base_counts.filter(estado_pago='PENDIENTE').count()
    
    # Ready for moodle: paid, partially collected or becado, and not matriculated yet.
    listos_moodle_count = base_counts.filter(
        estado_pago__in=Inscripcion.ESTADOS_PAGO_MATRICULABLES,
        estado_moodle='NO_MATRICULADO'
    ).count()
    todos_moodle_count = base_counts.filter(estado_moodle='NO_MATRICULADO').count()

    context = {
        'cursos': cursos,
        'inscripciones': inscripciones[:500],  # Limitar para rendimiento en UI
        'curso_id': curso_id,
        'estado_pago': estado_pago,
        'estado_moodle': estado_moodle,
        'search_query': search_query,
        # Indicadores
        'total_count': total_count,
        'pagados_count': pagados_count,
        'pendientes_count': pendientes_count,
        'listos_moodle_count': listos_moodle_count,
        'todos_moodle_count': todos_moodle_count,
        'active_tab': 'gestion',
    }
    return render(request, 'academico/gestion_inscripciones.html', context)


@login_required
def marcar_pago(request, inscripcion_id):
    if request.method == 'POST':
        inscripcion = get_object_or_404(Inscripcion, id=inscripcion_id)
        
        # Read the state parameter from POST (default: PAGADO)
        nuevo_estado = request.POST.get('estado_pago', 'PAGADO')
        inscripcion.estado_pago = nuevo_estado
        
        if nuevo_estado in Inscripcion.ESTADOS_PAGO_MATRICULABLES:
            inscripcion.fecha_pago = timezone.now().date()
        else:
            inscripcion.fecha_pago = None
        
        # Process optional file upload
        comprobante = request.FILES.get('comprobante_pago')
        if comprobante:
            inscripcion.comprobante_pago = comprobante
            
        inscripcion.save()
        messages.success(request, f'Se actualizó el estado de pago de {inscripcion.cursante.apellido}, {inscripcion.cursante.nombre} a "{inscripcion.get_estado_pago_display()}".')
    
    referer = request.META.get('HTTP_REFERER')
    if referer:
        return redirect(referer)
    return redirect('lista_inscripciones')


@login_required
def detalle_inscripcion(request, inscripcion_id):
    insc = get_object_or_404(Inscripcion, id=inscripcion_id)
    cursante = insc.cursante
    
    # Structured titles
    titulos = [{
        'codigo': t.codigo or '',
        'nombre': t.nombre,
        'nivel': t.nivel or ''
    } for t in cursante.titulos.all()]
    
    data = {
        'id': insc.id,
        'id_preinscripcion_syric': insc.id_preinscripcion_syric or '',
        'estado_pago': insc.estado_pago,
        'estado_pago_display': insc.get_estado_pago_display(),
        'fecha_pago': str(insc.fecha_pago) if insc.fecha_pago else '',
        'comprobante_url': insc.comprobante_pago.url if insc.comprobante_pago else '',
        'estado_moodle': insc.estado_moodle,
        'estado_moodle_display': insc.get_estado_moodle_display(),
        'estado_academico': insc.estado_academico,
        'estado_academico_display': insc.get_estado_academico_display(),
        'nota_final': str(insc.nota_final) if insc.nota_final is not None else '',
        'observaciones': insc.observaciones or '',
        
        'cursante': {
            'id': cursante.id,
            'dni': cursante.dni,
            'nombre': cursante.nombre,
            'apellido': cursante.apellido,
            'email': cursante.email or '',
            'telefono': cursante.telefono or '',
            'fecha_nacimiento': str(cursante.fecha_nacimiento) if cursante.fecha_nacimiento else '',
            'titulo_previo_raw': cursante.titulo_previo or '',
            'titulos': titulos
        },
        'curso': {
            'nombre': insc.curso.nombre,
            'id_syric': insc.curso.id_syric or '',
            'id_cohorte_moodle': insc.curso.id_cohorte_moodle or ''
        }
    }
    return JsonResponse(data)


@login_required
def guardar_observacion(request, inscripcion_id):
    if request.method == 'POST':
        insc = get_object_or_404(Inscripcion, id=inscripcion_id)
        insc.observaciones = request.POST.get('observaciones', '')
        insc.save()
        messages.success(request, f'Se actualizaron las observaciones de {insc.cursante.apellido}, {insc.cursante.nombre}.')
    
    referer = request.META.get('HTTP_REFERER')
    if referer:
        return redirect(referer)
    return redirect('lista_inscripciones')


@login_required
def descargar_csv_moodle(request, curso_id):
    incluir_todos = request.GET.get('todos') == '1'
    response, error = generate_moodle_csv_response(curso_id, incluir_todos=incluir_todos)
    if error:
        messages.error(request, error)
        return redirect('lista_inscripciones')
    return response


@login_required
def marcar_mail_enviado(request, inscripcion_id):
    if request.method == 'POST':
        inscripcion = get_object_or_404(Inscripcion, id=inscripcion_id)
        inscripcion.estado_moodle = 'MAIL_ENVIADO'
        inscripcion.save(update_fields=['estado_moodle'])
        messages.success(request, f'Se marcó el mail de credenciales como enviado para {inscripcion.cursante.apellido}, {inscripcion.cursante.nombre}.')
    
    referer = request.META.get('HTTP_REFERER')
    if referer:
        return redirect(referer)
    return redirect('lista_inscripciones')


# --- ETAPA 4: Portal del Docente y Cierre de Actas ---

@login_required
def portal_docente(request):
    # If superuser or staff, show all courses, otherwise show only assigned ones
    if request.user.is_superuser or request.user.is_staff:
        cursos = Curso.objects.filter(activo=True)
    else:
        cursos = Curso.objects.filter(docentes=request.user, activo=True)
        
    context = {
        'cursos': cursos,
        'active_tab': 'docente',
    }
    return render(request, 'academico/portal_docente.html', context)


@login_required
def acta_calificaciones(request, curso_id):
    curso = get_object_or_404(Curso, id=curso_id)
    
    # Check permissions
    if not (request.user.is_superuser or request.user.is_staff or request.user in curso.docentes.all()):
        messages.error(request, 'No tiene permisos para modificar las actas de este curso.')
        return redirect('portal_docente')
        
    inscripciones = Inscripcion.objects.filter(curso=curso, estado_moodle__in=['MATRICULADO', 'MAIL_ENVIADO']).select_related('cursante')

    if request.method == 'POST':
        try:
            with transaction.atomic():
                for insc in inscripciones:
                    id_str = str(insc.id)
                    estado_campo = f"estado_academico_{id_str}"
                    nota_campo = f"nota_final_{id_str}"
                    
                    if estado_campo in request.POST:
                        insc.estado_academico = request.POST.get(estado_campo)
                        
                    if nota_campo in request.POST:
                        nota_raw = request.POST.get(nota_campo, '').strip()
                        if nota_raw == '':
                            insc.nota_final = None
                        else:
                            # Clean decimal separators (replace comma with point)
                            nota_clean = nota_raw.replace(',', '.')
                            insc.nota_final = Decimal(nota_clean)
                            
                    insc.save()
            messages.success(request, 'Calificaciones y actas guardadas con éxito.')
            return redirect('portal_docente')
        except Exception as e:
            messages.error(request, f'Ocurrió un error al guardar las calificaciones: {str(e)}')

    context = {
        'curso': curso,
        'inscripciones': inscripciones,
        'active_tab': 'docente',
    }
    return render(request, 'academico/acta_calificaciones.html', context)


# --- ETAPA 4: Módulo de Certificación y Exportación ---

@login_required
def reporte_syric(request):
    if not (request.user.is_superuser or request.user.is_staff):
        messages.error(request, 'Acceso restringido únicamente para personal administrativo.')
        return redirect('portal_docente')
        
    cursos = Curso.objects.filter(activo=True)
    
    # Filters
    curso_id = request.GET.get('curso', '')
    certificado_estado = request.GET.get('certificado', '')
    
    inscripciones = Inscripcion.objects.all().select_related('cursante', 'curso')
    
    if curso_id:
        inscripciones = inscripciones.filter(curso_id=curso_id)
    if certificado_estado:
        inscripciones = inscripciones.filter(certificado_syric=certificado_estado)
        
    # approved and ready for certification counts
    approved_counts = Inscripcion.objects.filter(estado_academico='APROBADO')
    if curso_id:
        approved_counts = approved_counts.filter(curso_id=curso_id)
        
    total_aprobados = approved_counts.count()
    pendientes_syric = approved_counts.filter(certificado_syric='PENDIENTE').count()
    gestionados_syric = approved_counts.filter(certificado_syric='GESTIONADO').count()
    emitidos_syric = approved_counts.filter(certificado_syric='EMITIDO').count()

    context = {
        'cursos': cursos,
        'inscripciones': inscripciones[:500],
        'curso_id': curso_id,
        'certificado_estado': certificado_estado,
        # Counters
        'total_aprobados': total_aprobados,
        'pendientes_syric': pendientes_syric,
        'gestionados_syric': gestionados_syric,
        'emitidos_syric': emitidos_syric,
        'active_tab': 'certificacion',
    }
    return render(request, 'academico/reporte_syric.html', context)


@login_required
def descargar_aprobados_syric(request, curso_id):
    if not (request.user.is_superuser or request.user.is_staff):
        messages.error(request, 'Acceso denegado.')
        return redirect('portal_docente')
        
    curso = get_object_or_404(Curso, id=curso_id)
    
    # Select approved students
    aprobados = Inscripcion.objects.filter(
        curso=curso,
        estado_academico='APROBADO'
    ).select_related('cursante')
    
    if not aprobados.exists():
        messages.error(request, 'No hay cursantes aprobados en este curso para exportar.')
        return redirect('reporte_syric')
        
    # Create Excel sheet using openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Aprobados Syric"
    
    # Styling variables
    header_fill = PatternFill(start_color="3F51B5", end_color="3F51B5", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=11)
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    
    headers = [
        "DNI", 
        "Apellido", 
        "Nombre", 
        "Email", 
        "Teléfono", 
        "ID Preinscripción Syric", 
        "ID Capacitación", 
        "Nota Final", 
        "Estado Académico"
    ]
    
    # Write headers
    for col_idx, h_text in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=h_text)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center
        
    updated_ids = []
    
    # Write student details
    for row_idx, insc in enumerate(aprobados, 2):
        cursante = insc.cursante
        row_data = [
            cursante.dni,
            cursante.apellido,
            cursante.nombre,
            cursante.email or '',
            cursante.telefono or '',
            insc.id_preinscripcion_syric or '',
            curso.id_syric or '',
            float(insc.nota_final) if insc.nota_final is not None else '',
            insc.get_estado_academico_display()
        ]
        
        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = data_font
            if col_idx in [1, 6, 7, 8, 9]:
                cell.alignment = align_center
            else:
                cell.alignment = align_left
                
        updated_ids.append(insc.id)
        
    # Transaction atomic status updates
    with transaction.atomic():
        # Update state to GESTIONADO
        Inscripcion.objects.filter(id__in=updated_ids).update(certificado_syric='GESTIONADO')
        
    # Adjust column widths automatically
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
        
    # Prepare HTTP response
    fecha_str = timezone.now().strftime('%Y%m%d')
    codigo_curso = curso.id_syric or f"id_{curso.id}"
    filename = f"aprobados_syric_{codigo_curso}_{fecha_str}.xlsx"
    
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    
    return response


@login_required
def dashboard_principal(request):
    cursos_activos = Curso.objects.filter(activo=True)
    total_cursos = cursos_activos.count()
    total_cursantes = Cursante.objects.count()
    
    # Payments counts
    total_pagados = Inscripcion.objects.filter(estado_pago__in=Inscripcion.ESTADOS_PAGO_MATRICULABLES).count()
    total_pendientes = Inscripcion.objects.filter(estado_pago='PENDIENTE').count()
    
    # Moodle counts
    total_matriculados = Inscripcion.objects.filter(estado_moodle__in=['MATRICULADO', 'MAIL_ENVIADO']).count()
    
    # Academic performance counts
    total_aprobados = Inscripcion.objects.filter(estado_academico='APROBADO').count()
    total_desaprobados = Inscripcion.objects.filter(estado_academico='DESAPROBADO').count()
    total_libres = Inscripcion.objects.filter(estado_academico='LIBRE').count()
    total_cursando = Inscripcion.objects.filter(estado_academico='CURSANDO').count()
    
    # Course metrics summary table
    cursos_stats = []
    for curso in cursos_activos:
        inscs = Inscripcion.objects.filter(curso=curso)
        total_insc = inscs.count()
        
        pagados = inscs.filter(estado_pago__in=Inscripcion.ESTADOS_PAGO_MATRICULABLES).count()
        matriculados = inscs.filter(estado_moodle__in=['MATRICULADO', 'MAIL_ENVIADO']).count()
        aprobados = inscs.filter(estado_academico='APROBADO').count()
        desaprobados = inscs.filter(estado_academico='DESAPROBADO').count()
        libres = inscs.filter(estado_academico='LIBRE').count()
        
        # Calculate % Progress: (aprobados + desaprobados + libres) / total_insc * 100
        cerrados = aprobados + desaprobados + libres
        avance = round((cerrados / total_insc * 100), 2) if total_insc > 0 else 0.0
        
        # Docentes list
        docentes_list = [f"{doc.first_name} {doc.last_name}".strip() or doc.username for doc in curso.docentes.all()]
        docentes_str = ", ".join(docentes_list) if docentes_list else "Sin asignar"
        
        cursos_stats.append({
            'curso': curso,
            'docentes': docentes_str,
            'total_insc': total_insc,
            'pagados': pagados,
            'matriculados': matriculados,
            'aprobados': aprobados,
            'desaprobados': desaprobados,
            'libres': libres,
            'avance': avance
        })
        
    context = {
        'total_cursos': total_cursos,
        'total_cursantes': total_cursantes,
        'total_pagados': total_pagados,
        'total_pendientes': total_pendientes,
        'total_matriculados': total_matriculados,
        'total_aprobados': total_aprobados,
        'total_desaprobados': total_desaprobados,
        'total_libres': total_libres,
        'total_cursando': total_cursando,
        'cursos_stats': cursos_stats,
        'active_tab': 'dashboard',
    }
    return render(request, 'academico/dashboard.html', context)


@login_required
def editar_cursante(request, cursante_id):
    if request.method == 'POST':
        cursante = get_object_or_404(Cursante, id=cursante_id)
        cursante.nombre = request.POST.get('nombre', '').strip()
        cursante.apellido = request.POST.get('apellido', '').strip()
        cursante.dni = request.POST.get('dni', '').strip()
        cursante.email = request.POST.get('email', '').strip() or None
        cursante.telefono = request.POST.get('telefono', '').strip() or None
        cursante.save()
        messages.success(request, f'Ficha de {cursante.apellido}, {cursante.nombre} modificada con éxito.')
        
    referer = request.META.get('HTTP_REFERER')
    if referer:
        return redirect(referer)
    return redirect('lista_inscripciones')
